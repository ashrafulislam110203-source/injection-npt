"""Professional Excel workbooks built from the database."""
import datetime as dt
import io

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from database import db
from kpi import npt as npt_kpi, rejection as rej_kpi
from kpi.common import month_label

FONT = "Arial"
PCT = {"NPT %", "% of Total NPT", "% of Total Rejection", "Cumulative %"}
HOURS = ("(h)",)
HDR_FILL = PatternFill("solid", fgColor="1F3864")
THIN = Side(style="thin", color="BFBFBF")


def _fmt(col, series):
    if col in PCT:
        return "0.00%"
    if col.endswith("(h)") or col.endswith("(min)") or col.endswith("(kg)") or col.endswith("(ton)"):
        return "#,##0.00"
    if pd.api.types.is_float_dtype(series):
        return "#,##0.00"
    if pd.api.types.is_integer_dtype(series):
        return "#,##0"
    return None


def build_workbook(sheets, sum_cols=None, notes=None):
    """sheets: {name: DataFrame}. sum_cols: {name: [columns to total with SUBTOTAL]}. Returns bytes."""
    sum_cols = sum_cols or {}
    wb = Workbook()
    wb.remove(wb.active)
    for name, df in sheets.items():
        ws = wb.create_sheet(name[:31])
        if df is None or df.empty:
            ws["A1"] = "No data for this selection."
            ws["A1"].font = Font(name=FONT, italic=True)
            continue
        df = df.reset_index(drop=True)
        for j, col in enumerate(df.columns, 1):
            c = ws.cell(row=1, column=j, value=col)
            c.font = Font(name=FONT, bold=True, color="FFFFFF")
            c.fill = HDR_FILL
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        for j, col in enumerate(df.columns, 1):
            fmt = _fmt(col, df[col])
            width = max(len(str(col)) + 2, 10)
            for i, v in enumerate(df[col].tolist(), 2):
                if v is None or (isinstance(v, float) and np.isnan(v)) or v is pd.NaT:
                    v = None
                elif isinstance(v, pd.Timestamp):
                    v = v.to_pydatetime()
                elif isinstance(v, np.generic):
                    v = v.item()
                c = ws.cell(row=i, column=j, value=v)
                c.font = Font(name=FONT, size=10)
                if isinstance(v, dt.datetime):
                    c.number_format = "dd-mmm-yyyy hh:mm"
                elif isinstance(v, dt.date):
                    c.number_format = "dd-mmm-yyyy"
                elif fmt and isinstance(v, (int, float)):
                    c.number_format = fmt
                if v is not None:
                    width = max(width, min(len(str(v)) + 2, 60))
            ws.column_dimensions[get_column_letter(j)].width = width + (3 if isinstance(df[col].iloc[0], (dt.date, dt.datetime)) else 0)
        last = len(df) + 1
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{last}"
        cols = [c for c in sum_cols.get(name, []) if c in df.columns]
        if cols:
            tr = last + 1
            ws.cell(row=tr, column=1, value="Total").font = Font(name=FONT, bold=True)
            for c in cols:
                j = list(df.columns).index(c) + 1
                L = get_column_letter(j)
                cell = ws.cell(row=tr, column=j, value=f"=SUBTOTAL(109,{L}2:{L}{last})")
                cell.font = Font(name=FONT, bold=True)
                cell.number_format = _fmt(c, df[c]) or "#,##0"
                cell.border = Border(top=THIN)
    if notes:
        ws = wb.create_sheet("Notes")
        for i, line in enumerate(notes, 1):
            ws.cell(row=i, column=1, value=line).font = Font(name=FONT, size=10, bold=(i == 1))
        ws.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- friendly tables
def notes_for(settings):
    return [
        "How the numbers are calculated",
        "Source: NPT / Downtime report and Rejection report from Smart Manufacturing. Machine ID is normalised (IMM-380-03 = IMM-380-3).",
        "NPT (h): each downtime event is cut at midnight, so only the time inside each calendar day is counted.",
        "Occurrences = number of downtime events. Average / maximum duration are per event inside the selected period.",
        "NPT % = NPT hours / (machines x scheduled hours x days). Shown only when the number of machines is set in Settings.",
        "Rejection: report Quantity is in thousand pieces (x1000 = pieces). Report Weight is in tons (x1000 = kg).",
        "Rejection % is not calculated in this version because production quantity is not used.",
        "Causes come directly from the reports (trailing '*' and extra spaces removed). New causes are accepted automatically.",
    ]


def t_npt_events(nd):
    if nd.empty:
        return nd
    d = {"Date": nd.date, "Machine": nd.machine_id}
    if (nd["shift"] != "All").any():
        d["Shift"] = nd["shift"]
    d.update({"Cause": nd.cause, "NPT (h)": nd.npt_sec / 3600, "Event Start": nd.event_start, "Event End": nd.event_end})
    return pd.DataFrame(d)


def t_npt_pareto(nd):
    c = npt_kpi.by_cause(nd)
    if c.empty:
        return c
    return pd.DataFrame({"Cause": c.cause, "Occurrences": c.occurrences, "Total NPT (h)": c.total_h,
                         "% of Total NPT": c.pct, "Cumulative %": c.cum_pct, "Average Duration (min)": c.avg_min,
                         "Maximum Duration (min)": c.max_min})


def t_npt_machine(nd):
    m = npt_kpi.by_machine(nd)
    if m.empty:
        return m
    return pd.DataFrame({"Machine": m.machine_id, "Occurrences": m.occurrences, "Total NPT (h)": m.total_h,
                         "% of Total NPT": m.pct, "Average Duration (min)": m.avg_min,
                         "Maximum Duration (min)": m.max_min, "Top Cause": m.top_cause})


def t_npt_trend(nd, monthly=False, year=None):
    g = npt_kpi.by_month(nd, year) if monthly else npt_kpi.by_date(nd)
    if g.empty:
        return g
    first = g["month"].map(month_label) if monthly else g["date"]
    return pd.DataFrame({"Month" if monthly else "Date": first, "Occurrences": g.events,
                         "Machines Affected": g.machines, "NPT (h)": g.npt_h})


def t_rej_records(rj):
    if rj.empty:
        return rj
    return pd.DataFrame({"Date": rj.report_date, "Machine": rj.machine_id, "Item": rj.item_name, "WIP": rj.wip_name,
                         "Cause": rj.cause, "Rejected Pieces": rj.qty_pieces, "Rejected Weight (kg)": rj.weight_kg,
                         "Added By": rj.added_by, "Added Time": pd.to_datetime(rj.added_time)})


def _rej_table(g, key_label, extra=None):
    if g.empty:
        return g
    d = {key_label: g.iloc[:, 0], "Entries": g.entries, "Rejected Pieces": g.pieces, "Rejected Weight (kg)": g.kg,
         "% of Total Rejection": g.pct}
    if "cum_pct" in g and key_label == "Cause":
        d["Cumulative %"] = g.cum_pct
    if extra is not None and extra in g:
        d["Top Cause"] = g[extra]
    return pd.DataFrame(d)


def t_rej_pareto(rj):
    return _rej_table(rej_kpi.by_cause(rj), "Cause")


def t_rej_machine(rj):
    return _rej_table(rej_kpi.by_machine(rj), "Machine", "top_cause")


def t_rej_item(rj):
    return _rej_table(rej_kpi.by_item(rj), "Product")


def t_rej_trend(rj, monthly=False, year=None):
    g = rej_kpi.by_month(rj, year) if monthly else rej_kpi.by_date(rj)
    if g.empty:
        return g
    first = g["month"].map(month_label) if monthly else g["date"]
    return pd.DataFrame({"Month" if monthly else "Date": first, "Entries": g.entries, "Machines Affected": g.machines,
                         "Rejected Pieces": g.pieces, "Rejected Weight (kg)": g.kg})


def t_summary(title, start, end, ns, rs):
    rows = [("Period", f"{start:%d-%b-%Y} to {end:%d-%b-%Y}"),
            ("NPT", ""), ("Total NPT (h)", ns["npt_h"]), ("NPT events", ns["events"]), ("Machines with NPT", ns["machines"]),
            ("Average duration per event (min)", ns["avg_min"]), ("Longest event (min)", ns["max_min"]),
            ("Top NPT cause", ns["top_cause"]), ("NPT % of scheduled time", ns["npt_pct"]),
            ("Rejection", ""), ("Rejected pieces", rs["pieces"]), ("Rejected weight (kg)", rs["kg"]),
            ("Rejection entries", rs["entries"]), ("Machines with rejection", rs["machines"]),
            ("Top rejection cause", rs["top_cause"])]
    df = pd.DataFrame(rows, columns=["Item", "Value"])
    df["Value"] = df["Value"].map(lambda v: None if isinstance(v, float) and np.isnan(v) else (round(v, 4) if isinstance(v, float) else v))
    return df


def t_machine_summary(nd, rj):
    a = npt_kpi.by_machine(nd)[["machine_id", "occurrences", "total_h", "top_cause"]] if not nd.empty else pd.DataFrame(columns=["machine_id", "occurrences", "total_h", "top_cause"])
    b = rej_kpi.by_machine(rj)[["machine_id", "entries", "pieces", "kg", "top_cause"]] if not rj.empty else pd.DataFrame(columns=["machine_id", "entries", "pieces", "kg", "top_cause"])
    m = a.merge(b, how="outer", on="machine_id", suffixes=("_npt", "_rej")).sort_values("machine_id")
    if m.empty:
        return m
    return pd.DataFrame({"Machine": m.machine_id, "NPT Events": m.occurrences.fillna(0), "NPT (h)": m.total_h.fillna(0),
                         "Top NPT Cause": m.top_cause_npt, "Rejection Entries": m.entries.fillna(0),
                         "Rejected Pieces": m.pieces.fillna(0), "Rejected Weight (kg)": m.kg.fillna(0),
                         "Top Rejection Cause": m.top_cause_rej})


def report_workbook(start, end, settings=None, level="range", machines=None, causes=None, shift=None,
                    items=None, rej_causes=None, year=None):
    """level: daily | monthly | yearly | range. Everything is built from the database."""
    s = settings or db.get_settings()
    nd = npt_kpi.load_npt(start, end, machines, causes, shift, s)
    rj = rej_kpi.load_rejection(start, end, machines, items, rej_causes)
    days = (end - start).days + 1
    ns, rs = npt_kpi.summary(nd, s, days), rej_kpi.summary(rj, days)
    monthly = level == "yearly"
    title = {"daily": "Daily Summary", "monthly": "Monthly Summary", "yearly": "Yearly Summary"}.get(level, "Summary")
    sheets = {title: t_summary(title, start, end, ns, rs)}
    if level != "daily":
        sheets["Monthly NPT" if monthly else "Daily NPT"] = t_npt_trend(nd, monthly, year)
        sheets["Monthly Rejection" if monthly else "Daily Rejection"] = t_rej_trend(rj, monthly, year)
    sheets.update({"Machine Summary": t_machine_summary(nd, rj), "NPT Pareto": t_npt_pareto(nd), "NPT by Machine": t_npt_machine(nd),
                   "NPT Events": t_npt_events(nd), "Rejection Pareto": t_rej_pareto(rj), "Rejection by Machine": t_rej_machine(rj),
                   "Rejection by Product": t_rej_item(rj), "Rejection Records": t_rej_records(rj)})
    sums = {"NPT Pareto": ["Occurrences", "Total NPT (h)"], "NPT by Machine": ["Occurrences", "Total NPT (h)"],
            "NPT Events": ["NPT (h)"], "Rejection Pareto": ["Entries", "Rejected Pieces", "Rejected Weight (kg)"],
            "Rejection by Machine": ["Entries", "Rejected Pieces", "Rejected Weight (kg)"],
            "Rejection by Product": ["Entries", "Rejected Pieces", "Rejected Weight (kg)"],
            "Rejection Records": ["Rejected Pieces", "Rejected Weight (kg)"],
            "Machine Summary": ["NPT Events", "NPT (h)", "Rejection Entries", "Rejected Pieces", "Rejected Weight (kg)"],
            "Daily NPT": ["Occurrences", "NPT (h)"], "Monthly NPT": ["Occurrences", "NPT (h)"],
            "Daily Rejection": ["Entries", "Rejected Pieces", "Rejected Weight (kg)"],
            "Monthly Rejection": ["Entries", "Rejected Pieces", "Rejected Weight (kg)"]}
    return build_workbook(sheets, sums, notes_for(s))


def history_workbook(settings=None):
    """Everything stored, without internal database fields."""
    s = settings or db.get_settings()
    ev = db.query_df("SELECT machine_id, cause, start_time, end_time, duration_sec, added_by FROM npt_events ORDER BY start_time")
    if not ev.empty:
        ev = pd.DataFrame({"Machine": ev.machine_id, "Cause": ev.cause, "Event Start": pd.to_datetime(ev.start_time),
                           "Event End": pd.to_datetime(ev.end_time), "Duration (h)": ev.duration_sec / 3600, "Cause Added By": ev.added_by})
    rj = db.query_df("SELECT * FROM rejection ORDER BY report_date, added_time")
    if not rj.empty:
        rj["report_date"] = pd.to_datetime(rj.report_date).dt.date
    lo, hi = db.date_bounds()
    daily = pd.DataFrame()
    machine = pd.DataFrame()
    if lo:
        nd = npt_kpi.load_npt(lo, hi, settings=s)
        rjr = rej_kpi.load_rejection(lo, hi)
        daily = t_npt_trend(nd).merge(t_rej_trend(rjr), how="outer", on="Date", suffixes=(" (NPT)", " (Rejection)")) if not nd.empty and not rjr.empty else (t_npt_trend(nd) if not nd.empty else t_rej_trend(rjr))
        machine = t_machine_summary(nd, rjr)
    up = db.query_df("SELECT upload_time, report_date, report_type, file_name, records_total, records_valid, records_rejected, status, duplicate_status, error_count FROM upload_history ORDER BY id")
    up = up.rename(columns={"upload_time": "Upload Time", "report_date": "Report Date", "report_type": "Report Type",
                            "file_name": "File Name", "records_total": "Records Read", "records_valid": "Records Saved",
                            "records_rejected": "Rows Rejected", "status": "Status", "duplicate_status": "Duplicate Status",
                            "error_count": "Errors"})
    sheets = {"NPT Events": ev, "Rejection": t_rej_records(rj.rename(columns={}).assign(report_date=rj.report_date) if not rj.empty else rj),
              "Daily Summary": daily, "Machine Summary": machine, "Upload History": up}
    return build_workbook(sheets, notes=notes_for(s))
