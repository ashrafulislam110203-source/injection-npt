"""Tests that daily uploads are APPENDED, history is kept, month/year change works, and delete / reset / start date work.
Extra days are made by shifting the real 18 Sep sample files by N days (test data only, temporary database)."""
import datetime as dt, glob, io, os, re, sys, tempfile
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__))); sys.path.insert(0, ROOT)
folder = sys.argv[1] if len(sys.argv) > 1 else "/mnt/user-data/uploads"
os.environ["KPI_DB_PATH"] = os.path.join(tempfile.mkdtemp(), "append.db")
import openpyxl
from database import db, repository as repo
from processing.npt_processor import parse_npt
from processing.rejection_processor import parse_rejection
from kpi import npt, rejection
from utils import excel_export as ex
from ui.upload import _save_all

BASE = dt.date(2026, 9, 18)
f_npt = glob.glob(folder + "/Downtime*")[0]; f_rej = glob.glob(folder + "/Rejection*")[0]
tmp = tempfile.mkdtemp(); ok = []
def check(n, c, i=""): ok.append(bool(c)); print(("PASS  " if c else "FAIL  ") + n, i)
def count(t): return int(db.query_df(f"select count(*) c from {t}").c[0])

def shifted(src, days, kind):
    wb = openpyxl.load_workbook(src); ws = wb.active
    d = BASE + dt.timedelta(days=days)
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, dt.datetime):
                c.value = c.value + dt.timedelta(days=days)
            elif isinstance(c.value, str) and " To " in c.value and re.search(r"\d+/\d+/\d{4}", c.value):
                c.value = f"{d:%d/%m/%Y} To {d:%d/%m/%Y}" if kind == "npt" else f"{d.month}/{d.day}/{d.year} To {d.month}/{d.day}/{d.year}"
    p = os.path.join(tmp, f"{kind}_{d}.xlsx"); wb.save(p); return p, d

db.init_db(); S = db.get_settings()
check("database starts empty (real data starts 1 Sep 2026)", count("npt_events") == 0 and count("rejection") == 0 and S["data_start_date"] == "2026-09-01")
days = [0, 1, 2, 13, 43]     # 18 Sep, 19 Sep, 20 Sep, 1 Oct, 31 Oct  -> crosses month end
per_day = {}
for n in days:
    pn, d = shifted(f_npt, n, "npt"); pr, _ = shifted(f_rej, n, "rej")
    a = parse_npt(pn, "n", d, S); b = parse_rejection(pr, "r", d, S)
    pend = dict(reports={"NPT": a, "Rejection": b}, errors={}, dups={"NPT": repo.existing_info(a), "Rejection": repo.existing_info(b)}, date=d)
    res = _save_all(pend, {"NPT": "Skip Duplicate Records", "Rejection": "Skip Duplicate Records"})
    per_day[d] = (a.rows_valid, b.rows_valid)
    print(d, res)
check("every day appended, none overwritten: rejection rows", count("rejection") == 61 * len(days), count("rejection"))
n_ev = count("npt_events")
check("NPT events appended (no double for carry-over events)", n_ev > 195 * (len(days) - 2), n_ev)
lo, hi = db.date_bounds(); check("date range covers all days", lo == BASE and hi >= BASE + dt.timedelta(days=43), (lo, hi))
d1 = rejection.load_rejection(BASE, BASE).qty_pieces.sum(); d2 = rejection.load_rejection(BASE + dt.timedelta(days=1), BASE + dt.timedelta(days=1)).qty_pieces.sum()
check("each day keeps its own totals", round(d1) == 2541 and round(d2) == 2541, (d1, d2))
sep = rejection.load_rejection(dt.date(2026, 9, 1), dt.date(2026, 9, 30)); octo = rejection.load_rejection(dt.date(2026, 10, 1), dt.date(2026, 10, 31))
check("monthly totals: Sep = 3 days, Oct = 2 days", round(sep.qty_pieces.sum()) == 3 * 2541 and round(octo.qty_pieces.sum()) == 2 * 2541)
ym = rejection.by_month(rejection.load_rejection(dt.date(2026, 1, 1), dt.date(2026, 12, 31)), 2026)
check("yearly view has 12 months, Sep and Oct filled", len(ym) == 12 and ym.set_index("month").loc["2026-09", "pieces"] > 0 and ym.set_index("month").loc["2026-10", "pieces"] > 0 and ym.set_index("month").loc["2026-01", "pieces"] == 0)
nm = npt.by_month(npt.load_npt(dt.date(2026, 1, 1), dt.date(2026, 12, 31), settings=S), 2026)
check("NPT monthly has Sep and Oct", nm.set_index("month").loc["2026-09", "npt_h"] > 0 and nm.set_index("month").loc["2026-10", "npt_h"] > 0)
y = ex.report_workbook(dt.date(2026, 1, 1), dt.date(2026, 12, 31), level="yearly", year=2026)
wb = openpyxl.load_workbook(io.BytesIO(y)); check("yearly export: 12 month rows + total", wb["Monthly NPT"].max_row == 14, wb["Monthly NPT"].max_row)
h = openpyxl.load_workbook(io.BytesIO(ex.history_workbook())); check("history export contains all days", h["Rejection"].max_row - 2 == 61 * len(days))

# re-upload same day: replace keeps others
pn, d = shifted(f_npt, 1, "npt"); a = parse_npt(pn, "n", d, S)
pend = dict(reports={"NPT": a}, errors={}, dups={"NPT": repo.existing_info(a)}, date=d)
before = count("npt_events"); _save_all(pend, {"NPT": "Replace Existing Data"})
check("replace one day leaves other days untouched", count("npt_events") == before, (before, count("npt_events")))

# delete + reset
n = repo.delete_day("Rejection", BASE + dt.timedelta(days=2))
check("delete one day only", n == 61 and count("rejection") == 61 * (len(days) - 1), n)
hist = db.query_df("select status from upload_history where status='Deleted'"); check("delete is recorded in upload history", len(hist) == 1)
path = repo.reset_all()
check("reset makes a backup first and empties data", os.path.exists(path) and count("npt_events") == 0 and count("rejection") == 0, path)
check("upload history kept after reset", count("upload_history") > 5)
print("\nALL PASSED" if all(ok) else "\nSOME FAILED")
